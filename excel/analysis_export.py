"""Create a polished Excel workbook from the cleaned retail datasets.

The workbook is intentionally secondary to SQL and Power BI. It contains
analysis-ready summary tables and two simple charts, not a second data model.
"""

from __future__ import annotations

import argparse
from pathlib import Path
from typing import Iterable

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.worksheet.table import Table, TableStyleInfo


NAVY = "17324D"
BLUE = "2F75B5"
TEAL = "2A9D8F"
GOLD = "F4B942"
PALE_BLUE = "EAF2F8"
LIGHT_GREY = "E7E9ED"
WHITE = "FFFFFF"
TEXT = "243447"
MONEY_FORMAT = '$#,##0.00;[Red]-$#,##0.00'
PERCENT_FORMAT = "0.00%"
INTEGER_FORMAT = "#,##0"
DECIMAL_FORMAT = "#,##0.00"


def _safe_divide(numerator: float, denominator: float) -> float:
    return float(numerator / denominator) if denominator else 0.0


def _load_cleaned(cleaned_dir: Path) -> dict[str, pd.DataFrame]:
    names = ["customers", "products", "stores", "sales", "inventory", "returns"]
    missing = [name for name in names if not (cleaned_dir / f"{name}.csv").exists()]
    if missing:
        raise FileNotFoundError(f"Missing cleaned CSV files: {', '.join(missing)}")
    return {name: pd.read_csv(cleaned_dir / f"{name}.csv") for name in names}


def _build_analysis(tables: dict[str, pd.DataFrame]) -> dict[str, object]:
    sales = tables["sales"].copy()
    products = tables["products"].copy()
    stores = tables["stores"].copy()
    customers = tables["customers"].copy()
    inventory = tables["inventory"].copy()
    returns = tables["returns"].copy()
    sales["order_date"] = pd.to_datetime(sales["order_date"])
    inventory["date"] = pd.to_datetime(inventory["date"])

    monthly = (
        sales.assign(month=sales["order_date"].dt.strftime("%Y-%m"))
        .groupby("month", as_index=False)
        .agg(
            revenue=("net_revenue", "sum"),
            profit=("profit", "sum"),
            orders=("order_id", "nunique"),
            units_sold=("quantity", "sum"),
        )
    )
    monthly["previous_month_revenue"] = monthly["revenue"].shift(1)
    monthly["month_over_month_growth"] = (
        (monthly["revenue"] - monthly["previous_month_revenue"])
        / monthly["previous_month_revenue"].replace(0, np.nan)
    )

    store_performance = (
        sales.merge(stores, on="store_id", how="left", validate="many_to_one")
        .groupby(["store_id", "store_name", "city", "province", "store_type"], as_index=False)
        .agg(
            revenue=("net_revenue", "sum"),
            profit=("profit", "sum"),
            orders=("order_id", "nunique"),
            units_sold=("quantity", "sum"),
        )
    )
    store_performance["profit_margin"] = store_performance["profit"] / store_performance["revenue"]
    store_performance["average_order_value"] = store_performance["revenue"] / store_performance["orders"]
    store_performance = store_performance.sort_values("revenue", ascending=False).reset_index(drop=True)
    store_performance.insert(0, "revenue_rank", np.arange(1, len(store_performance) + 1))

    product_performance = (
        sales.merge(products[["product_id", "product_name", "category", "subcategory"]], on="product_id", how="left", validate="many_to_one")
        .groupby(["product_id", "product_name", "category", "subcategory"], as_index=False)
        .agg(
            revenue=("net_revenue", "sum"),
            profit=("profit", "sum"),
            units_sold=("quantity", "sum"),
            sales_lines=("order_id", "size"),
            average_discount=("discount_pct", "mean"),
        )
    )
    product_performance["profit_margin"] = product_performance["profit"] / product_performance["revenue"]
    product_performance = product_performance.sort_values("revenue", ascending=False).reset_index(drop=True)
    product_performance.insert(0, "revenue_rank", np.arange(1, len(product_performance) + 1))

    customer_analysis = (
        sales.loc[sales["customer_id"] != "CUST_UNKNOWN"]
        .merge(customers[["customer_id", "customer_name", "city", "province", "customer_segment"]], on="customer_id", how="left", validate="many_to_one")
        .groupby(["customer_id", "customer_name", "city", "province", "customer_segment"], as_index=False)
        .agg(
            orders=("order_id", "nunique"),
            total_spend=("net_revenue", "sum"),
            gross_profit=("profit", "sum"),
            first_order=("order_date", "min"),
            latest_order=("order_date", "max"),
        )
    )
    customer_analysis["average_order_value"] = customer_analysis["total_spend"] / customer_analysis["orders"]
    customer_analysis["repeat_customer"] = np.where(customer_analysis["orders"] >= 2, "Yes", "No")
    customer_analysis["first_order"] = customer_analysis["first_order"].dt.strftime("%Y-%m-%d")
    customer_analysis["latest_order"] = customer_analysis["latest_order"].dt.strftime("%Y-%m-%d")
    customer_analysis = customer_analysis.sort_values("total_spend", ascending=False).reset_index(drop=True)

    inventory_analysis = (
        inventory.merge(products[["product_id", "product_name", "category", "unit_cost"]], on="product_id", how="left", validate="many_to_one")
        .groupby(["product_id", "product_name", "category"], as_index=False)
        .agg(
            inventory_observations=("date", "size"),
            average_closing_stock=("closing_stock", "mean"),
            total_units_sold=("units_sold", "sum"),
            stockout_observations=("closing_stock", lambda values: int((values == 0).sum())),
            below_reorder_observations=("closing_stock", "size"),
        )
    )
    below_reorder_counts = (
        inventory.assign(below=inventory["closing_stock"] < inventory["reorder_level"])
        .groupby("product_id")["below"]
        .sum()
    )
    inventory_analysis["below_reorder_observations"] = inventory_analysis["product_id"].map(below_reorder_counts).astype(int)
    inventory_analysis["stockout_rate"] = inventory_analysis["stockout_observations"] / inventory_analysis["inventory_observations"]
    inventory_analysis = inventory_analysis.sort_values(["total_units_sold", "average_closing_stock"], ascending=[True, False]).reset_index(drop=True)

    sold_by_category = (
        sales.merge(products[["product_id", "category"]], on="product_id", how="left", validate="many_to_one")
        .groupby("category", as_index=False)["quantity"]
        .sum()
        .rename(columns={"quantity": "units_sold"})
    )
    returned_by_category = (
        returns.merge(products[["product_id", "category"]], on="product_id", how="left", validate="many_to_one")
        .groupby("category", as_index=False)["quantity_returned"]
        .sum()
        .rename(columns={"quantity_returned": "units_returned"})
    )
    returns_category = sold_by_category.merge(returned_by_category, on="category", how="left")
    returns_category["units_returned"] = returns_category["units_returned"].fillna(0).astype(int)
    returns_category["return_rate"] = returns_category["units_returned"] / returns_category["units_sold"]
    returns_category = returns_category.sort_values("return_rate", ascending=False).reset_index(drop=True)
    returns_reason = (
        returns.groupby("return_reason", as_index=False)
        .agg(return_records=("return_id", "size"), units_returned=("quantity_returned", "sum"))
        .sort_values("units_returned", ascending=False)
        .reset_index(drop=True)
    )
    returns_reason["share_of_returned_units"] = returns_reason["units_returned"] / returns_reason["units_returned"].sum()

    inventory_value = inventory.merge(products[["product_id", "unit_cost"]], on="product_id", how="left", validate="many_to_one")
    monthly_inventory_value = (inventory_value["closing_stock"] * inventory_value["unit_cost"]).groupby(inventory_value["date"]).sum()
    metrics = {
        "Total Revenue": float(sales["net_revenue"].sum()),
        "Gross Profit": float(sales["profit"].sum()),
        "Profit Margin": _safe_divide(sales["profit"].sum(), sales["net_revenue"].sum()),
        "Orders": int(sales["order_id"].nunique()),
        "Average Order Value": _safe_divide(sales["net_revenue"].sum(), sales["order_id"].nunique()),
        "Units Sold": int(sales["quantity"].sum()),
        "Customers": int(sales.loc[sales["customer_id"] != "CUST_UNKNOWN", "customer_id"].nunique()),
        "Return Rate": _safe_divide(returns["quantity_returned"].sum(), sales["quantity"].sum()),
        "Stockout Rate": float((inventory["closing_stock"] == 0).mean()),
        "Inventory Turnover": _safe_divide(sales["cost"].sum(), monthly_inventory_value.mean()),
    }
    return {
        "metrics": metrics,
        "monthly": monthly,
        "stores": store_performance,
        "products": product_performance,
        "customers": customer_analysis,
        "inventory": inventory_analysis,
        "returns_category": returns_category,
        "returns_reason": returns_reason,
    }


def _clean_header(value: object) -> str:
    return str(value).replace("_", " ").title()


def _write_dataframe(
    worksheet,
    frame: pd.DataFrame,
    *,
    start_row: int,
    table_name: str,
    title: str | None = None,
) -> tuple[int, int]:
    data = frame.copy()
    data.columns = [_clean_header(column) for column in data.columns]
    if title:
        worksheet.cell(start_row, 1, title)
        worksheet.cell(start_row, 1).font = Font(name="Aptos Display", size=13, bold=True, color=NAVY)
        header_row = start_row + 1
    else:
        header_row = start_row

    for row in dataframe_to_rows(data, index=False, header=True):
        worksheet.append(row)
    end_row = header_row + len(data)
    end_col = len(data.columns)
    table = Table(
        displayName=table_name,
        ref=f"A{header_row}:{get_column_letter(end_col)}{end_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
    return end_row, end_col


def _format_sheet(worksheet, *, header_row: int = 4) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = f"A{header_row + 1}"
    worksheet.auto_filter.ref = worksheet.dimensions
    thin = Side(style="thin", color=LIGHT_GREY)
    for cell in worksheet[header_row]:
        if cell.value is not None:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=thin)
    worksheet.row_dimensions[header_row].height = 24

    for column_cells in worksheet.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        worksheet.column_dimensions[letter].width = min(max(max_length + 2, 11), 34)

    headers = {cell.column: str(cell.value).lower() for cell in worksheet[header_row] if cell.value}
    for column_index, header in headers.items():
        letter = get_column_letter(column_index)
        cells = worksheet[f"{letter}{header_row + 1}:{letter}{worksheet.max_row}"]
        number_format = None
        if any(term in header for term in ["revenue", "profit", "spend", "value", "cost"]):
            number_format = MONEY_FORMAT
        if any(term in header for term in ["margin", "rate", "growth", "discount", "share"]):
            number_format = PERCENT_FORMAT
        if any(term in header for term in ["orders", "units", "rank", "frequency", "observations", "sales lines"]):
            number_format = INTEGER_FORMAT
        if "average closing" in header:
            number_format = DECIMAL_FORMAT
        if number_format:
            for row in cells:
                row[0].number_format = number_format


def _add_sheet_title(worksheet, title: str, note: str, width: int) -> None:
    end_letter = get_column_letter(max(width, 6))
    worksheet.merge_cells(f"A1:{end_letter}1")
    worksheet["A1"] = title
    worksheet["A1"].font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    worksheet["A1"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 30
    worksheet.merge_cells(f"A2:{end_letter}2")
    worksheet["A2"] = note
    worksheet["A2"].font = Font(name="Aptos", size=10, italic=True, color=TEXT)
    worksheet["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    worksheet.row_dimensions[2].height = 30


def _create_data_sheet(
    workbook: Workbook,
    name: str,
    title: str,
    note: str,
    frame: pd.DataFrame,
    table_name: str,
) -> None:
    worksheet = workbook.create_sheet(name)
    for row in dataframe_to_rows(frame.rename(columns=_clean_header), index=False, header=True):
        worksheet.append(row)
    worksheet.insert_rows(1, amount=3)
    _add_sheet_title(worksheet, title, note, len(frame.columns))
    end_row = 4 + len(frame)
    end_col = len(frame.columns)
    table = Table(displayName=table_name, ref=f"A4:{get_column_letter(end_col)}{end_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False, showColumnStripes=False)
    worksheet.add_table(table)
    _format_sheet(worksheet, header_row=4)


def _style_executive_summary(worksheet, metrics: dict[str, float | int]) -> None:
    worksheet.sheet_view.showGridLines = False
    worksheet.merge_cells("A1:J2")
    worksheet["A1"] = "Retail Operations Analytics — Executive Summary"
    worksheet["A1"].font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    worksheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    worksheet["A1"].alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 30
    worksheet.row_dimensions[2].height = 14
    worksheet["A3"] = "Synthetic portfolio data | Detailed tables are available on the remaining sheets."
    worksheet["A3"].font = Font(name="Aptos", size=10, italic=True, color=TEXT)

    worksheet["A5"] = "KPI"
    worksheet["B5"] = "Value"
    for cell in worksheet[5][:2]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
    for row_number, (label, value) in enumerate(metrics.items(), start=6):
        worksheet.cell(row_number, 1, label)
        worksheet.cell(row_number, 2, value)
        worksheet.cell(row_number, 1).font = Font(name="Aptos", bold=True, color=TEXT)
        worksheet.cell(row_number, 2).font = Font(name="Aptos", size=11, bold=True, color=NAVY)
        worksheet.cell(row_number, 1).fill = PatternFill("solid", fgColor=PALE_BLUE if row_number % 2 == 0 else WHITE)
        worksheet.cell(row_number, 2).fill = PatternFill("solid", fgColor=PALE_BLUE if row_number % 2 == 0 else WHITE)
        if label in {"Total Revenue", "Gross Profit", "Average Order Value"}:
            worksheet.cell(row_number, 2).number_format = MONEY_FORMAT
        elif "Rate" in label or "Margin" in label:
            worksheet.cell(row_number, 2).number_format = PERCENT_FORMAT
        elif label == "Inventory Turnover":
            worksheet.cell(row_number, 2).number_format = '0.00x'
        else:
            worksheet.cell(row_number, 2).number_format = INTEGER_FORMAT

    worksheet.column_dimensions["A"].width = 25
    worksheet.column_dimensions["B"].width = 18
    for letter in "CDEFGHIJ":
        worksheet.column_dimensions[letter].width = 13


def export_analysis_workbook(
    cleaned_dir: Path | str,
    output_path: Path | str,
) -> Path:
    """Build, save, and structurally verify the Excel analysis workbook."""
    cleaned_path = Path(cleaned_dir)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    analysis = _build_analysis(_load_cleaned(cleaned_path))

    workbook = Workbook()
    executive = workbook.active
    executive.title = "Executive Summary"
    _style_executive_summary(executive, analysis["metrics"])

    _create_data_sheet(
        workbook,
        "Monthly Sales",
        "Monthly Sales",
        "One row per calendar month. Growth compares each complete month with the prior month.",
        analysis["monthly"],
        "MonthlySalesTable",
    )
    _create_data_sheet(
        workbook,
        "Store Performance",
        "Store Performance",
        "Revenue, profit, margin, orders, AOV, and units at store grain.",
        analysis["stores"],
        "StorePerformanceTable",
    )
    _create_data_sheet(
        workbook,
        "Product Performance",
        "Product Performance",
        "Products are ranked by net revenue; margin is profit divided by revenue.",
        analysis["products"],
        "ProductPerformanceTable",
    )
    _create_data_sheet(
        workbook,
        "Customer Analysis",
        "Customer Analysis",
        "One row per purchasing customer. CUST_UNKNOWN is excluded from customer behavior metrics.",
        analysis["customers"],
        "CustomerAnalysisTable",
    )
    _create_data_sheet(
        workbook,
        "Inventory Analysis",
        "Inventory Analysis",
        "Monthly store-product observations summarized to product grain.",
        analysis["inventory"],
        "InventoryAnalysisTable",
    )

    returns_sheet = workbook.create_sheet("Returns Analysis")
    _add_sheet_title(
        returns_sheet,
        "Returns Analysis",
        "Unit return rates use returned units divided by sold units. Tables are aggregated separately before comparison.",
        6,
    )
    end_category, _ = _write_dataframe(
        returns_sheet,
        analysis["returns_category"],
        start_row=4,
        table_name="ReturnsCategoryTable",
        title="Returns by Category",
    )
    _write_dataframe(
        returns_sheet,
        analysis["returns_reason"],
        start_row=end_category + 3,
        table_name="ReturnsReasonTable",
        title="Return Reasons",
    )
    returns_sheet.sheet_view.showGridLines = False
    returns_sheet.freeze_panes = "A6"
    for row in [5, end_category + 4]:
        for cell in returns_sheet[row]:
            if cell.value is not None:
                cell.fill = PatternFill("solid", fgColor=NAVY)
                cell.font = Font(name="Aptos", bold=True, color=WHITE)
    for column_cells in returns_sheet.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in column_cells)
        returns_sheet.column_dimensions[letter].width = min(max(max_length + 2, 12), 32)
    for row in returns_sheet.iter_rows():
        for cell in row:
            if cell.row > 5 and cell.column in [4]:
                cell.number_format = PERCENT_FORMAT

    monthly_sheet = workbook["Monthly Sales"]
    line_chart = LineChart()
    line_chart.title = "Monthly Revenue Trend"
    line_chart.style = 13
    line_chart.y_axis.title = "Revenue (CAD)"
    line_chart.x_axis.title = "Month"
    line_chart.y_axis.numFmt = '$#,##0'
    line_chart.height = 8
    line_chart.width = 16
    line_chart.add_data(Reference(monthly_sheet, min_col=2, min_row=4, max_row=monthly_sheet.max_row), titles_from_data=True)
    line_chart.set_categories(Reference(monthly_sheet, min_col=1, min_row=5, max_row=monthly_sheet.max_row))
    line_chart.legend = None
    executive.add_chart(line_chart, "D5")

    store_sheet = workbook["Store Performance"]
    bar_chart = BarChart()
    bar_chart.type = "bar"
    bar_chart.style = 10
    bar_chart.title = "Revenue by Store"
    bar_chart.x_axis.title = "Revenue (CAD)"
    bar_chart.y_axis.title = "Store"
    bar_chart.x_axis.numFmt = '$#,##0'
    bar_chart.height = 8
    bar_chart.width = 16
    revenue_column = list(analysis["stores"].columns).index("revenue") + 1
    store_name_column = list(analysis["stores"].columns).index("store_name") + 1
    bar_chart.add_data(Reference(store_sheet, min_col=revenue_column, min_row=4, max_row=store_sheet.max_row), titles_from_data=True)
    bar_chart.set_categories(Reference(store_sheet, min_col=store_name_column, min_row=5, max_row=store_sheet.max_row))
    bar_chart.legend = None
    bar_chart.dLbls = DataLabelList()
    bar_chart.dLbls.showVal = False
    executive.add_chart(bar_chart, "D21")

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(destination)

    check = load_workbook(destination, read_only=False, data_only=False)
    expected_sheets = [
        "Executive Summary", "Monthly Sales", "Store Performance", "Product Performance",
        "Customer Analysis", "Inventory Analysis", "Returns Analysis",
    ]
    if check.sheetnames != expected_sheets:
        raise RuntimeError(f"Workbook sheet verification failed: {check.sheetnames}")
    if len(check["Executive Summary"]._charts) != 2:
        raise RuntimeError("Workbook chart verification failed")
    for name in expected_sheets[1:]:
        if check[name].max_row < 5:
            raise RuntimeError(f"Workbook sheet is unexpectedly empty: {name}")
    check.close()

    print("Excel analysis workbook created successfully.")
    print(f"  sheets     : {len(expected_sheets)}")
    print("  charts     : 2")
    print(f"  workbook   : {destination.resolve()}")
    return destination


def parse_args() -> argparse.Namespace:
    project_root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--cleaned-dir", type=Path, default=project_root / "data" / "cleaned")
    parser.add_argument("--output", type=Path, default=project_root / "excel" / "retail_operations_analysis.xlsx")
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    export_analysis_workbook(args.cleaned_dir, args.output)

